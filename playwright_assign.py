import os
import re
import json
import time
from datetime import datetime

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font
from playwright.sync_api import sync_playwright


# ============================================================
# CONFIGURATION
# ============================================================

CHROME_PATH = r"C:\Program Files\Google\Chrome\Application\chrome.exe"

PROJECT_DIR = os.getcwd()

PROFILE_DIR = os.path.join(PROJECT_DIR, "whatsapp_profile")
SCREENSHOT_DIR = os.path.join(PROJECT_DIR, "screenshots")
REPORT_DIR = os.path.join(PROJECT_DIR, "reports")

CONTACTS_FILE = os.path.join(PROJECT_DIR, "mycontacts.xlsx")

JSON_REPORT = os.path.join(REPORT_DIR, "whatsapp_report.json")
EXCEL_REPORT = os.path.join(REPORT_DIR, "whatsapp_report.xlsx")

WHATSAPP_URL = "https://web.whatsapp.com"


# Create required folders
os.makedirs(SCREENSHOT_DIR, exist_ok=True)
os.makedirs(REPORT_DIR, exist_ok=True)


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def clean_phone_number(phone):
    """
    Convert phone number into digits only.

    Examples:
        +91 9876543210 -> 919876543210
        +91-9876543210 -> 919876543210
    """

    phone_text = str(phone).strip()

    # Remove Excel numeric formatting such as 919876543210.0
    if phone_text.endswith(".0"):
        phone_text = phone_text[:-2]

    # Keep digits only
    phone_number = re.sub(r"\D", "", phone_text)

    return phone_number


def safe_filename(text):
    """
    Make text safe to use as a Windows filename.
    """

    text = str(text)

    safe = "".join(
        character
        for character in text
        if character.isalnum() or character in (" ", "_", "-")
    )

    return safe.strip()


def save_reports(results):
    """
    Save the collected results to both JSON and Excel.
    """

    # --------------------------------------------------------
    # JSON REPORT
    # --------------------------------------------------------

    with open(JSON_REPORT, "w", encoding="utf-8") as file:
        json.dump(
            results,
            file,
            indent=4,
            ensure_ascii=False
        )

    # --------------------------------------------------------
    # EXCEL REPORT
    # --------------------------------------------------------

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "WhatsApp Report"

    headers = [
        "Name",
        "Phone",
        "Message Sent",
        "Status",
        "Screenshot",
        "Last Message 1",
        "Last Message 2",
        "Last Message 3",
        "Timestamp",
        "Error"
    ]

    worksheet.append(headers)

    # Make headers bold
    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    # Add results
    for result in results:

        last_messages = result.get("last_3_messages", [])

        message_1 = last_messages[0] if len(last_messages) > 0 else ""
        message_2 = last_messages[1] if len(last_messages) > 1 else ""
        message_3 = last_messages[2] if len(last_messages) > 2 else ""

        worksheet.append([
            result.get("name", ""),
            result.get("phone", ""),
            result.get("message_sent", ""),
            result.get("status", ""),
            result.get("screenshot", ""),
            message_1,
            message_2,
            message_3,
            result.get("timestamp", ""),
            result.get("error", "")
        ])

    worksheet.freeze_panes = "A2"

    # Adjust column widths
    for column in worksheet.columns:

        max_length = 0
        column_letter = column[0].column_letter

        for cell in column:

            if cell.value is not None:
                cell_length = len(str(cell.value))

                if cell_length > max_length:
                    max_length = cell_length

        worksheet.column_dimensions[column_letter].width = min(
            max_length + 2,
            60
        )

    workbook.save(EXCEL_REPORT)


def extract_last_three_messages(page):
    """
    Extract the last 3 messages visible in the current WhatsApp chat.
    """

    messages = page.locator('div[data-testid="msg-container"]')

    count = messages.count()

    print(f"Messages found: {count}")

    extracted_messages = []

    # Work backwards so we get the newest messages first
    for index in range(count - 1, -1, -1):

        try:
            message_text = messages.nth(index).inner_text().strip()

            if message_text:
                extracted_messages.append(message_text)

            if len(extracted_messages) == 3:
                break

        except Exception:
            continue

    # Reverse so messages are oldest -> newest
    extracted_messages.reverse()

    return extracted_messages


def read_contacts():
    """
    Read mycontacts.xlsx.

    Expected columns:

        Name
        Phone
        Message
    """

    if not os.path.exists(CONTACTS_FILE):
        raise FileNotFoundError(
            f"mycontacts.xlsx was not found:\n{CONTACTS_FILE}"
        )

    workbook = load_workbook(
        CONTACTS_FILE,
        data_only=True
    )

    worksheet = workbook.active

    contacts = []

    # Read rows starting from row 2
    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=2,
            values_only=True
        ),
        start=2
    ):

        name = row[0] if len(row) > 0 else None
        phone = row[1] if len(row) > 1 else None
        message_template = row[2] if len(row) > 2 else None

        # Ignore completely empty rows
        if not name and not phone:
            continue

        if not phone:
            print(
                f"Skipping row {row_number}: phone number is missing."
            )
            continue

        contacts.append({
            "name": str(name).strip() if name else "there",
            "phone": phone,
            "message_template": (
                str(message_template).strip()
                if message_template
                else ""
            )
        })

    return contacts


# ============================================================
# MAIN BOT
# ============================================================

def main():

    print("=" * 60)
    print("WHATSAPP MESSAGE SENDER + SMART DATA EXTRACTOR")
    print("=" * 60)

    # --------------------------------------------------------
    # READ EXCEL
    # --------------------------------------------------------

    contacts = read_contacts()

    print(f"\nContacts loaded: {len(contacts)}")

    if not contacts:
        print("No contacts found in mycontacts.xlsx.")
        return

    for index, contact in enumerate(contacts, start=1):
        print(
            f"{index}. {contact['name']} - "
            f"{contact['phone']}"
        )

    print("\nStarting WhatsApp automation...")

    results = []

    # --------------------------------------------------------
    # START PLAYWRIGHT
    # --------------------------------------------------------

    with sync_playwright() as playwright:

        print("\nOpening Chrome...")

        context = playwright.chromium.launch_persistent_context(
            user_data_dir=PROFILE_DIR,
            headless=False,
            executable_path=CHROME_PATH,
            viewport=None
        )

        # Reuse existing page if available
        if context.pages:
            page = context.pages[0]
        else:
            page = context.new_page()

        # ----------------------------------------------------
        # OPEN WHATSAPP WEB
        # ----------------------------------------------------

        print("Opening WhatsApp Web...")

        page.goto(
            WHATSAPP_URL,
            wait_until="domcontentloaded",
            timeout=60000
        )

        print("\nWaiting for WhatsApp Web...")
        page.wait_for_timeout(8000)

        print("WhatsApp Web loaded.")

        # ----------------------------------------------------
        # PROCESS EACH CONTACT
        # ----------------------------------------------------

        for index, contact in enumerate(contacts, start=1):

            name = contact["name"]
            phone = clean_phone_number(contact["phone"])
            message_template = contact["message_template"]

            print("\n" + "=" * 60)
            print(
                f"Processing {index}/{len(contacts)}: "
                f"{name}"
            )
            print(f"Phone: {phone}")
            print("=" * 60)

            timestamp = datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            )

            result = {
                "name": name,
                "phone": phone,
                "message_sent": "",
                "status": "Failed",
                "screenshot": "",
                "last_3_messages": [],
                "timestamp": timestamp,
                "error": ""
            }

            try:

                # ------------------------------------------------
                # CREATE PERSONALIZED MESSAGE
                # ------------------------------------------------

                if message_template:

                    personalized_message = (
                        message_template.replace(
                            "{name}",
                            name
                        )
                    )

                else:

                    personalized_message = (
                        f"Hello {name}, "
                        f"this is a test message."
                    )

                print(
                    f"Message: {personalized_message}"
                )

                # ------------------------------------------------
                # OPEN CHAT
                # ------------------------------------------------

                chat_url = (
                    "https://web.whatsapp.com/send"
                    f"?phone={phone}"
                )

                print("Opening chat...")

                page.goto(
                    chat_url,
                    wait_until="domcontentloaded",
                    timeout=60000
                )

                # Give WhatsApp time to load the chat
                page.wait_for_timeout(5000)

                # ------------------------------------------------
                # FIND MESSAGE BOX
                # ------------------------------------------------

                print("Finding message box...")

                message_box = page.locator(
                    'div[contenteditable="true"]'
                ).last

                message_box.wait_for(
                    state="visible",
                    timeout=30000
                )

                # ------------------------------------------------
                # SEND MESSAGE
                # ------------------------------------------------

                print("Sending message...")

                message_box.fill(
                    personalized_message
                )

                message_box.press("Enter")

                # Wait for WhatsApp to send it
                page.wait_for_timeout(3000)

                print("Message sent.")

                result["message_sent"] = (
                    personalized_message
                )

                # ------------------------------------------------
                # SCREENSHOT
                # ------------------------------------------------

                filename = (
                    f"{index:02d}_"
                    f"{safe_filename(name)}_"
                    f"{phone}.png"
                )

                screenshot_path = os.path.join(
                    SCREENSHOT_DIR,
                    filename
                )

                print("Taking screenshot...")

                page.screenshot(
                    path=screenshot_path,
                    full_page=False
                )

                result["screenshot"] = os.path.join(
                    "screenshots",
                    filename
                )

                print(
                    f"Screenshot saved: {screenshot_path}"
                )

                # ------------------------------------------------
                # EXTRACT LAST 3 MESSAGES
                # ------------------------------------------------

                print("Extracting last 3 messages...")

                last_3_messages = (
                    extract_last_three_messages(page)
                )

                result["last_3_messages"] = (
                    last_3_messages
                )

                print("\nLast 3 messages:")

                for message_number, message in enumerate(
                    last_3_messages,
                    start=1
                ):
                    print(
                        f"\nMessage {message_number}:"
                    )
                    print(message)

                # ------------------------------------------------
                # SUCCESS
                # ------------------------------------------------

                result["status"] = "Success"

                print(
                    f"\n✓ {name} processed successfully."
                )

            except Exception as error:

                result["error"] = str(error)

                print(
                    f"\n✗ Failed to process {name}"
                )

                print(
                    f"Error: {error}"
                )

            # Add result
            results.append(result)

            # Save reports after every contact
            # so successful results aren't lost if
            # a later contact fails.
            save_reports(results)

            print("\nReports updated.")

            # Small pause before next contact
            page.wait_for_timeout(2000)

        # --------------------------------------------------------
        # FINISHED
        # --------------------------------------------------------

        print("\n" + "=" * 60)
        print("AUTOMATION COMPLETE")
        print("=" * 60)

        successful = sum(
            1
            for result in results
            if result["status"] == "Success"
        )

        failed = len(results) - successful

        print(f"\nTotal contacts : {len(results)}")
        print(f"Successful     : {successful}")
        print(f"Failed         : {failed}")

        print("\nReports:")
        print(f"JSON  : {JSON_REPORT}")
        print(f"Excel : {EXCEL_REPORT}")

        print("\nScreenshots:")
        print(SCREENSHOT_DIR)

        print("\nBrowser will remain open.")
        input(
            "\nPress ENTER to close WhatsApp..."
        )

        context.close()


# ============================================================
# RUN PROGRAM
# ============================================================

if __name__ == "__main__":
    main()