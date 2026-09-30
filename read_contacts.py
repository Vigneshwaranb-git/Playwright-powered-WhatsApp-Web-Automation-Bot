from openpyxl import load_workbook

FILE_NAME = "mycontacts.xlsx"

workbook = load_workbook(FILE_NAME)
sheet = workbook.active

print("Reading contacts...")
print()

for row in sheet.iter_rows(min_row=2, values_only=True):

    name = row[0]
    phone = row[1]
    message = row[2]

    print("Name   :", name)
    print("Phone  :", phone)
    print("Message:", message)
    print("-" * 40)